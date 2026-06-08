#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
将 docs/答辩PPT.md 转换为 PowerPoint (.pptx) 文件
"""

import os
import re
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ============================================
# 配置
# ============================================
INPUT_MD = r"C:\Users\xiexq\campus_swap\docs\答辩PPT.md"
OUTPUT_PPTX = r"C:\Users\xiexq\campus_swap\docs\校园闲置物品交换系统_答辩PPT.pptx"

# 主题色
THEME_PRIMARY = RGBColor(0x1E, 0x3A, 0x5F)      # 深蓝主色
THEME_ACCENT = RGBColor(0x2E, 0x8B, 0x57)       # 绿色强调
THEME_BG = RGBColor(0xFA, 0xFA, 0xFA)          # 浅灰背景
THEME_TEXT = RGBColor(0x33, 0x33, 0x33)        # 深灰文字
THEME_LIGHT_TEXT = RGBColor(0x66, 0x66, 0x66)  # 浅灰文字

TITLE_FONT_SIZE = 32
SUBTITLE_FONT_SIZE = 20
BODY_FONT_SIZE = 14
SMALL_FONT_SIZE = 12
HEADER_FONT_SIZE = 16


# ============================================
# 工具函数
# ============================================
def add_title_shape(slide, text, top, left, width, height, font_size=TITLE_FONT_SIZE, bold=True,
                    color=THEME_PRIMARY, align=PP_ALIGN.LEFT):
    """添加标题文本框"""
    shape = slide.shapes.add_textbox(left, top, width, height)
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Microsoft YaHei"
    p.alignment = align
    return shape


def add_body_text(slide, text, top, left, width, height, font_size=BODY_FONT_SIZE,
                  color=THEME_TEXT, bold=False, align=PP_ALIGN.LEFT):
    """添加正文文本框"""
    shape = slide.shapes.add_textbox(left, top, width, height)
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Microsoft YaHei"
    p.alignment = align
    return shape


def add_paragraphs(slide, lines, top, left, width, height, font_size=BODY_FONT_SIZE,
                   color=THEME_TEXT, bullet=False):
    """添加多段落文本"""
    shape = slide.shapes.add_textbox(left, top, width, height)
    tf = shape.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = "Microsoft YaHei"
        p.level = 0 if not bullet else 0
        if bullet:
            p.text = "•  " + line
    return shape


def add_decorative_bar(slide, top, left, width, height, color):
    """添加装饰条"""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def parse_markdown_table(lines):
    """解析Markdown表格"""
    if not lines:
        return [], []
    # 过滤掉空行和分隔线
    rows = []
    for line in lines:
        line = line.strip()
        if line and not line.startswith('| ---'):
            # 解析单元格
            cells = [c.strip() for c in line.split('|')]
            cells = [c for c in cells if c]  # 过滤空项
            if cells:
                rows.append(cells)
    if len(rows) == 0:
        return [], []
    headers = rows[0]
    data = rows[1:] if len(rows) > 1 else []
    return headers, data


def add_table(slide, headers, data, top, left, width, row_height=Inches(0.4), font_size=12):
    """添加表格"""
    if not headers:
        return None
    num_rows = 1 + len(data)
    num_cols = len(headers)
    table = slide.shapes.add_table(num_rows, num_cols, left, top, width,
                                   Inches(row_height.inches * num_rows)).table

    # 设置表头
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = THEME_PRIMARY
        para = cell.text_frame.paragraphs[0]
        para.font.size = Pt(font_size)
        para.font.bold = True
        para.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        para.font.name = "Microsoft YaHei"
        para.alignment = PP_ALIGN.CENTER

    # 设置数据行
    for row_idx, row_data in enumerate(data):
        for col_idx, cell_text in enumerate(row_data):
            if col_idx < num_cols:
                cell = table.cell(row_idx + 1, col_idx)
                cell.text = str(cell_text)
                if row_idx % 2 == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(0xF0, 0xF4, 0xF8)
                para = cell.text_frame.paragraphs[0]
                para.font.size = Pt(font_size)
                para.font.color.rgb = THEME_TEXT
                para.font.name = "Microsoft YaHei"
                para.alignment = PP_ALIGN.LEFT

    return table


# ============================================
# 解析Markdown并生成幻灯片
# ============================================
def parse_md_to_slides(md_path):
    """将Markdown拆分为幻灯片"""
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    # 按 --- 分隔
    raw_slides = re.split(r'\n---\s*\n', content)
    slides_data = []

    for raw in raw_slides:
        raw = raw.strip()
        if not raw:
            continue

        lines = raw.split('\n')
        title = ""
        content_lines = []
        tables = []
        in_table = False
        table_lines = []

        for line in lines:
            stripped = line.strip()
            if not stripped:
                if in_table and table_lines:
                    tables.append(parse_markdown_table(table_lines))
                    table_lines = []
                    in_table = False
                continue

            # 标题检测
            if stripped.startswith('# ') and not title:
                title = stripped[2:].strip()
            elif stripped.startswith('## ') and not title:
                title = stripped[3:].strip()
            elif stripped.startswith('### '):
                content_lines.append({"type": "subtitle", "text": stripped[4:].strip()})
            elif stripped.startswith('|') and '|' in stripped[1:]:
                in_table = True
                table_lines.append(stripped)
            elif stripped.startswith('```'):
                pass  # 忽略代码块标记，把代码当普通文本处理
            elif stripped.startswith('> '):
                content_lines.append({"type": "quote", "text": stripped[2:].strip()})
            elif stripped.startswith('- ') or stripped.startswith('* '):
                content_lines.append({"type": "bullet", "text": stripped[2:].strip()})
            elif re.match(r'^\d+\.', stripped):
                content_lines.append({"type": "numbered", "text": stripped.strip()})
            elif stripped.startswith('**') and stripped.endswith('**'):
                content_lines.append({"type": "bold", "text": stripped.strip('*').strip()})
            else:
                content_lines.append({"type": "text", "text": stripped.strip()})

        if in_table and table_lines:
            tables.append(parse_markdown_table(table_lines))

        slides_data.append({
            "title": title,
            "content": content_lines,
            "tables": tables
        })

    return slides_data


# ============================================
# 创建幻灯片
# ============================================
def create_presentation(slides_data, output_path):
    """创建PPTX文件"""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 使用空白布局
    blank_layout = prs.slide_layouts[6]

    for idx, slide_data in enumerate(slides_data):
        slide = prs.slides.add_slide(blank_layout)
        title = slide_data.get("title", "")
        content = slide_data.get("content", [])
        tables = slide_data.get("tables", [])

        is_cover = idx == 0 or (idx == len(slides_data) - 1 and "感谢" in title)

        # 背景装饰
        # 顶部色条
        add_decorative_bar(slide, Inches(0), Inches(0), prs.slide_width, Inches(0.15), THEME_PRIMARY)
        # 左侧色条
        add_decorative_bar(slide, Inches(0), Inches(0), Inches(0.05), prs.slide_height, THEME_PRIMARY)

        if is_cover:
            # 封面样式
            # 大标题居中偏上
            add_title_shape(
                slide, title,
                Inches(2.5), Inches(1.5), Inches(10.333), Inches(1.5),
                font_size=44, bold=True, color=THEME_PRIMARY, align=PP_ALIGN.CENTER
            )

            # 副标题
            subtitle_lines = [c["text"] for c in content if c.get("type") in ("text", "subtitle", "quote")]
            if subtitle_lines:
                y_pos = Inches(4.2)
                for line in subtitle_lines:
                    add_body_text(
                        slide, line,
                        y_pos, Inches(2.5), Inches(8.333), Inches(0.5),
                        font_size=20, color=THEME_LIGHT_TEXT, align=PP_ALIGN.CENTER
                    )
                    y_pos += Inches(0.5)
        else:
            # 普通页面样式
            # 页面标题
            add_title_shape(
                slide, title,
                Inches(0.5), Inches(0.4), Inches(11), Inches(0.9),
                font_size=TITLE_FONT_SIZE, bold=True, color=THEME_PRIMARY
            )

            # 分隔线
            add_decorative_bar(slide, Inches(1.15), Inches(0.4), Inches(2), Inches(0.03), THEME_ACCENT)

            # 正文内容区域
            y_pos = Inches(1.4)
            x_pos = Inches(0.5)
            content_width = Inches(11.8)

            for item in content:
                item_type = item.get("type", "text")
                text = item.get("text", "")
                if not text:
                    continue

                if item_type == "subtitle":
                    add_body_text(
                        slide, text,
                        y_pos, x_pos, content_width, Inches(0.4),
                        font_size=SUBTITLE_FONT_SIZE, color=THEME_ACCENT, bold=True
                    )
                    y_pos += Inches(0.45)
                elif item_type == "bullet":
                    # 处理粗体标记 **xxx**
                    is_bold = '**' in text
                    clean_text = text.replace('**', '')
                    add_body_text(
                        slide, "• " + clean_text,
                        y_pos, x_pos + Inches(0.2), content_width - Inches(0.2), Inches(0.35),
                        font_size=BODY_FONT_SIZE, color=THEME_TEXT, bold=is_bold
                    )
                    y_pos += Inches(0.38)
                elif item_type == "numbered":
                    add_body_text(
                        slide, text,
                        y_pos, x_pos + Inches(0.2), content_width - Inches(0.2), Inches(0.35),
                        font_size=BODY_FONT_SIZE, color=THEME_TEXT
                    )
                    y_pos += Inches(0.38)
                elif item_type == "quote":
                    add_body_text(
                        slide, text,
                        y_pos, x_pos + Inches(0.3), content_width - Inches(0.3), Inches(0.35),
                        font_size=SMALL_FONT_SIZE, color=THEME_LIGHT_TEXT
                    )
                    y_pos += Inches(0.35)
                elif item_type == "bold":
                    add_body_text(
                        slide, text,
                        y_pos, x_pos, content_width, Inches(0.35),
                        font_size=BODY_FONT_SIZE, color=THEME_TEXT, bold=True
                    )
                    y_pos += Inches(0.38)
                else:
                    # 普通文本
                    add_body_text(
                        slide, text,
                        y_pos, x_pos, content_width, Inches(0.35),
                        font_size=BODY_FONT_SIZE, color=THEME_TEXT
                    )
                    y_pos += Inches(0.38)

            # 添加表格
            for headers, data in tables:
                if headers:
                    table = add_table(
                        slide, headers, data,
                        y_pos, x_pos, content_width,
                        row_height=Inches(0.35), font_size=11
                    )
                    if table:
                        y_pos += Inches(0.35 * (1 + len(data)))

    # 保存
    prs.save(output_path)
    print(f"✅ PPT已生成: {output_path}")
    print(f"   共 {len(slides_data)} 页")
    return output_path


# ============================================
# 主程序
# ============================================
if __name__ == "__main__":
    try:
        slides_data = parse_md_to_slides(INPUT_MD)
        print(f"[INFO] 从 Markdown 解析到 {len(slides_data)} 页幻灯片")
        create_presentation(slides_data, OUTPUT_PPTX)
        print(f"[DONE] 完成！请查看: {OUTPUT_PPTX}")
    except Exception as e:
        print(f"[ERROR] 错误: {e}")
        import traceback
        traceback.print_exc()
